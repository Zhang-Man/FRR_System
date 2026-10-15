# -*- coding: utf-8 -*-
"""Bulk data loading from record.xlsx and API verification tests."""
import os
import pytest


@pytest.fixture(scope='module')
def record_xlsx_path():
    """Get path to record.xlsx if it exists."""
    paths = [
        os.path.join(os.path.dirname(__file__), '..', 'record', 'record.xlsx'),
        os.path.join(os.path.dirname(__file__), '..', 'record', 'record.csv'),
    ]
    for p in paths:
        if os.path.exists(p):
            return p
    return None


def _build_app_and_client():
    """Create a fresh app and client with test DB (module-scoped helper)."""
    import tempfile
    fd, db_path = tempfile.mkstemp(suffix='.db', prefix='frr_load_')
    os.close(fd)
    os.environ['DB_PATH'] = db_path
    os.environ['SECRET_KEY'] = 'test-load-secret'
    os.environ['LOG_LEVEL'] = 'WARNING'

    from app import create_app
    app = create_app()
    app.config['TESTING'] = True
    client = app.test_client()

    return client, db_path


@pytest.fixture(scope='module')
def load_client():
    """Module-scoped client for loading tests (shared DB)."""
    client, db_path = _build_app_and_client()
    yield client
    try:
        os.unlink(db_path)
    except (OSError, PermissionError):
        pass


@pytest.fixture(scope='module')
def load_handler_team():
    """handler operations team mapping."""
    # 示例处理人小组成员映射。请替换为贵机构真实的姓名/工号后再运行。
    return {
        '张三': 'S001', '李四': 'S002', '王五': 'S003',
        '赵六': 'S004', '钱七': 'S005', '孙八': 'S006',
        '周九': 'S007', '吴十': 'S008', '郑一': 'S009',
        '冯二': 'S010',
    }


def test_import_team_members(load_client, load_handler_team, record_xlsx_path):
    """Import handler team members into DB."""
    if record_xlsx_path is None:
        pytest.skip('record.xlsx not found')

    # Login admin
    resp = load_client.post('/api/auth/login', json={'user_id': 'FrrSystemAdmin'})
    token = resp.get_json()['token']

    # Import all team members as handlers
    users = [{'user_id': uid, 'name': name, 'role': 'handler'}
             for name, uid in load_handler_team.items()]
    resp = load_client.post('/api/users/import', json={'users': users}, headers={
        'Authorization': f'Bearer {token}'
    })
    assert resp.get_json()['imported'] == len(load_handler_team)


def test_load_tickets_from_excel(load_client, load_handler_team, record_xlsx_path):
    """Load tickets from record.xlsx into the database."""
    if record_xlsx_path is None:
        pytest.skip('record.xlsx not found')

    import openpyxl
    import sqlite3
    import os as _os

    # Login admin
    resp = load_client.post('/api/auth/login', json={'user_id': 'FrrSystemAdmin'})
    token = resp.get_json()['token']

    # Check for a submitter user
    resp = load_client.get('/api/users', headers={
        'Authorization': f'Bearer {token}'
    })
    users = resp.get_json()
    submitter = next((u for u in users if u['role'] == 'submitter'), None)
    if submitter is None:
        # Import a submitter
        load_client.post('/api/users/import', json={
            'users': [{'user_id': 'T001', 'name': '张老师', 'role': 'submitter'}]
        }, headers={'Authorization': f'Bearer {token}'})
        submitter = {'user_id': 'T001', 'name': '张老师'}

    submitter_id = submitter['user_id']
    submitter_name = submitter['name']

    wb = openpyxl.load_workbook(record_xlsx_path)
    db_path = _os.environ['DB_PATH']
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row

    TEAM = load_handler_team
    DEFAULT_HANDLER_ID = 'S001'

    def guess_category(text):
        text = str(text) if text else ''
        if '网络' in text or '上网' in text or '网线' in text or 'IP' in text.upper():
            if '慢' in text:
                return '网速慢'
            return '网络不通'
        if '打印' in text or '设备' in text or '电脑' in text or '机器' in text:
            return '设备故障'
        if '系统' in text or '安装' in text or '软件' in text:
            return '设备故障'
        return '其他'

    def match_assignee(persons_raw):
        if not persons_raw:
            return DEFAULT_HANDLER_ID, '张三'
        raw = persons_raw.replace('、', ' ').replace('，', ' ').replace(',', ' ')
        names = [n.strip() for n in raw.split() if n.strip() and len(n.strip()) >= 2]
        for name in names:
            if name in TEAM:
                return TEAM[name], name
            for team_name, team_id in TEAM.items():
                if name in team_name or team_name in name:
                    return team_id, team_name
        return DEFAULT_HANDLER_ID, '张三'

    ticket_count = 0
    for sname in wb.sheetnames[1:]:
        ws = wb[sname]
        hdr = [str(c.value or '') for c in ws[2]] if ws.max_row > 2 else [str(c.value or '') for c in ws[1]]

        time_col = reason_col = solution_col = person_col = 0
        for i, h in enumerate(hdr):
            if '时间' in h: time_col = i
            if '原因' in h or '描述' in h: reason_col = i
            if '处理' in h or '解决' in h: solution_col = i
            if '维修' in h or '维' in h: person_col = i

        for row in ws.iter_rows(min_row=3, values_only=True):
            if not row[reason_col]: continue

            time_raw = str(row[time_col] or '')
            reason = str(row[reason_col] or '')[:100]
            solution = str(row[solution_col] or '')[:200] if row[solution_col] else ''
            persons_raw = str(row[person_col] or '') if row[person_col] else ''

            created_at = time_raw.replace('.', '-')
            if not created_at or created_at == 'None' or len(created_at) < 7:
                continue
            parts = created_at.split('-')
            if len(parts) >= 3:
                created_at = f'{parts[0]}-{parts[1].zfill(2)}-{parts[2].zfill(2)}'
            created_at += ' 09:00:00' if ' ' not in created_at else ''

            category = guess_category(reason)
            location = reason[:30] if reason else '未知'
            assignee_id, assignee_name = match_assignee(persons_raw)

            if solution:
                status = 'resolved'
                completed_at = created_at.replace('09:00:00', '17:00:00')
            else:
                status = 'processing'
                completed_at = None

            conn.execute('''
                INSERT INTO tickets
                (reporter_id, reporter_name, contact, location, category, description,
                 assignee_id, assignee_name, status, solution, created_at, updated_at, completed_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ''', (
                submitter_id, submitter_name, '13800001111', location, category, reason,
                assignee_id, assignee_name, status, solution,
                created_at, created_at, completed_at
            ))
            ticket_count += 1

    conn.commit()

    total = conn.execute("SELECT COUNT(*) FROM tickets").fetchone()[0]
    pending = conn.execute("SELECT COUNT(*) FROM tickets WHERE status='pending'").fetchone()[0]
    processing = conn.execute("SELECT COUNT(*) FROM tickets WHERE status='processing'").fetchone()[0]
    resolved = conn.execute("SELECT COUNT(*) FROM tickets WHERE status='resolved'").fetchone()[0]
    conn.close()

    assert ticket_count > 0, 'No tickets loaded from Excel'
    assert total == ticket_count
    assert processing + resolved == ticket_count
    print(f'\n   Loaded {ticket_count} tickets (pending={pending}, processing={processing}, resolved={resolved})')


def test_api_verification_after_load(load_client, record_xlsx_path):
    """API checks after data load."""
    if record_xlsx_path is None:
        pytest.skip('record.xlsx not found')

    resp = load_client.post('/api/auth/login', json={'user_id': 'FrrSystemAdmin'})
    token = resp.get_json()['token']

    # Admin sees all
    resp = load_client.get('/api/tickets', headers={
        'Authorization': f'Bearer {token}'
    })
    assert len(resp.get_json()) > 100

    # Filter pending
    resp = load_client.get('/api/tickets?status=pending', headers={
        'Authorization': f'Bearer {token}'
    })
    assert len(resp.get_json()) >= 0

    # Filter resolved
    resp = load_client.get('/api/tickets?status=resolved', headers={
        'Authorization': f'Bearer {token}'
    })
    assert len(resp.get_json()) >= 0

    # CSV export
    resp = load_client.get('/api/tickets/export', headers={
        'Authorization': f'Bearer {token}'
    })
    assert resp.status_code == 200
    assert resp.data[:3] == b'\xef\xbb\xbf'

    # submitter login
    resp = load_client.post('/api/auth/login', json={'user_id': 'T001'})
    t_token = resp.get_json()['token']
    resp = load_client.get('/api/tickets', headers={
        'Authorization': f'Bearer {t_token}'
    })
    assert len(resp.get_json()) > 0

    # handler login
    resp = load_client.post('/api/auth/login', json={'user_id': 'S001'})
    s_token = resp.get_json()['token']
    resp = load_client.get('/api/tickets', headers={
        'Authorization': f'Bearer {s_token}'
    })
    assert len(resp.get_json()) >= 0
